"""Build the Northwind Robotics demo model: a small three-statement model of a
fictional company, laid out the way analysts like it, four quarters then the
year (1Q 2Q 3Q 4Q FY), with the house colour conventions (blue inputs, black
formulas, green links to other sheets).

    uv run --with openpyxl demo/northwind.py demo/northwind.xlsx

`build()` takes a few knobs so `agent.py` can play an AI agent revising it.
"""

from __future__ import annotations

import os
import sys
import tempfile
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

YEARS = [2024, 2025, 2026, 2027]
ACTUAL_YEARS = {2024, 2025}
FIRST = 3  # first period column (C)

NAVY = PatternFill("solid", fgColor="1F3A5F")
INPUT = PatternFill("solid", fgColor="FFF2CC")
BAND = PatternFill("solid", fgColor="F2F5F9")
WHITE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BLUE = Font(name="Calibri", size=11, color="1F4FD1")
BLACK = Font(name="Calibri", size=11, color="000000")
GREEN = Font(name="Calibri", size=11, color="0E7C3A")
BOLD = Font(name="Calibri", size=11, bold=True)
TITLE = Font(name="Calibri", size=16, bold=True, color="1F3A5F")
NOTE = Font(name="Calibri", size=10, italic=True, color="6B7280")
MEMO = Font(name="Calibri", size=10, italic=True, color="4B5563")
TOP = Border(top=Side(style="thin", color="1F3A5F"))
TOTAL = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

NUM = '#,##0.0;(#,##0.0);"–"'
PCT = '0.0%;(0.0%);"–"'
MULT = '0.0"x"'


@dataclass
class Knobs:
    """What the 'agent' changes between versions."""

    unit_growth: list[float] = field(default_factory=lambda: [0.22, 0.18])  # 2026, 2027 YoY
    asp_change: list[float] = field(default_factory=lambda: [0.03, 0.02])
    gross_margin: list[float] = field(default_factory=lambda: [0.46, 0.47])
    services: bool = False  # add a Services revenue line
    wacc: float = 0.095
    terminal_growth: float = 0.03


@dataclass
class Period:
    year: int
    q: int  # 0 = full year

    @property
    def actual(self) -> bool:
        return self.year in ACTUAL_YEARS

    @property
    def label(self) -> str:
        suffix = "A" if self.actual else "E"
        return f"FY{str(self.year)[2:]}{suffix}" if self.q == 0 else f"{self.q}Q{str(self.year)[2:]}{suffix}"


PERIODS = [Period(y, q) for y in YEARS for q in (1, 2, 3, 4, 0)]
COL = {(p.year, p.q): FIRST + i for i, p in enumerate(PERIODS)}


def L(col: int) -> str:
    return get_column_letter(col)


def quarter_cols(year: int) -> list[str]:
    return [L(COL[(year, q)]) for q in (1, 2, 3, 4)]


def prev_quarter(p: Period) -> Period | None:
    if p.q == 1:
        return Period(p.year - 1, 4) if p.year - 1 in YEARS else None
    return Period(p.year, p.q - 1)


# Reported history for the actual quarters (units in thousands, ASP in USD k).
UNITS = {2024: [3.10, 3.42, 3.71, 4.25], 2025: [4.02, 4.38, 4.61, 5.20]}
ASP = {2024: [41.0, 41.6, 42.1, 42.4], 2025: [43.0, 43.3, 43.9, 44.6]}
GM = {2024: [0.402, 0.411, 0.418, 0.426], 2025: [0.431, 0.437, 0.444, 0.452]}
SEASONAL = [0.95, 1.0, 1.02, 1.12]


class Sheet:
    def __init__(self, wb: Workbook, title: str, subtitle: str):
        self.ws = wb.create_sheet(title)
        self.row = 5
        ws = self.ws
        ws["A1"] = f"Northwind Robotics — {title}"
        ws["A1"].font = TITLE
        ws["A2"] = subtitle
        ws["A2"].font = NOTE
        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 9
        for p in PERIODS:
            c = COL[(p.year, p.q)]
            ws.column_dimensions[L(c)].width = 11 if p.q else 12
            cell = ws.cell(4, c, p.label)
            cell.font = WHITE
            cell.fill = NAVY
            cell.alignment = Alignment(horizontal="right")
        for c, text in ((1, "Line item"), (2, "Unit")):
            cell = ws.cell(4, c, text)
            cell.font = WHITE
            cell.fill = NAVY
        ws.freeze_panes = "C5"
        ws.sheet_view.showGridLines = False

    def section(self, text: str) -> None:
        self.row += 1
        self.ws.cell(self.row, 1, text).font = BOLD
        self.row += 1

    def line(self, label, unit, fn, *, fmt=NUM, font=None, indent=0, total=False, memo=False, annual="sum"):
        """Write one row. `fn(p)` returns a value or a formula for a quarter;
        annual columns follow `annual`: 'sum', 'last', a callable, or None."""
        ws, r = self.ws, self.row
        ws.cell(r, 1, label).alignment = Alignment(indent=indent)
        ws.cell(r, 2, unit).font = NOTE
        if memo:
            ws.cell(r, 1).font = MEMO
        if total:
            ws.cell(r, 1).font = BOLD
        for p in PERIODS:
            c = COL[(p.year, p.q)]
            if p.q:
                v = fn(p)
            elif annual == "sum":
                v = f"=SUM({quarter_cols(p.year)[0]}{r}:{quarter_cols(p.year)[3]}{r})"
            elif annual == "last":
                v = f"={quarter_cols(p.year)[3]}{r}"
            elif callable(annual):
                v = annual(p)
            else:
                v = None
            if v is None:
                continue
            cell = ws.cell(r, c, v)
            cell.number_format = fmt
            is_input = not (isinstance(v, str) and v.startswith("="))
            cell.font = font or (BLUE if is_input else (GREEN if isinstance(v, str) and "!" in v else BLACK))
            if memo:
                cell.font = Font(name="Calibri", size=10, italic=True, color=cell.font.color.rgb if cell.font.color else "000000")
            if is_input and not p.actual:
                cell.fill = INPUT
            if total:
                cell.font = Font(name="Calibri", size=11, bold=True, color=cell.font.color.rgb if cell.font.color else "000000")
                cell.border = TOTAL
            if not p.q:
                cell.fill = cell.fill if cell.fill.fgColor.rgb not in (None, "00000000") else BAND
        self.row += 1
        return r

    def blank(self) -> None:
        self.row += 1


def ref(sheet: str, col: int, row: int) -> str:
    return f"'{sheet}'!{L(col)}{row}"


def build(k: Knobs) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------ assumptions
    a = Sheet(wb, "Assumptions", "Inputs in blue on yellow; history in blue. USD millions unless stated.")
    a.section("Volume and price")
    units = a.line(
        "Robots shipped",
        "k",
        lambda p: UNITS[p.year][p.q - 1] if p.actual else f"={L(COL[(p.year - 1, p.q)])}{a.row}*(1+{L(COL[(p.year, p.q)])}{a.row + 1})",
        fmt='#,##0.00',
    )
    a.line("Unit growth, year on year", "%", lambda p: None if p.actual else k.unit_growth[p.year - 2026], fmt=PCT, indent=1, annual=None)
    asp = a.line(
        "Average selling price",
        "USD k",
        lambda p: ASP[p.year][p.q - 1] if p.actual else f"={L(COL[(p.year - 1, p.q)])}{a.row}*(1+{L(COL[(p.year, p.q)])}{a.row + 1})",
        fmt='#,##0.0',
        annual=lambda p: f"=SUMPRODUCT({quarter_cols(p.year)[0]}{units}:{quarter_cols(p.year)[3]}{units},{quarter_cols(p.year)[0]}{a.row}:{quarter_cols(p.year)[3]}{a.row})/{L(COL[(p.year, 0)])}{units}",
    )
    a.line("Price change, year on year", "%", lambda p: None if p.actual else k.asp_change[p.year - 2026], fmt=PCT, indent=1, annual=None)
    if k.services:
        a.line("Services attach rate", "% of hw", lambda p: 0.0 if p.actual else 0.06 + 0.01 * (p.year - 2026), fmt=PCT, annual=None)
    attach = a.row - 1 if k.services else None
    a.section("Margins and costs")
    gm = a.line("Gross margin", "%", lambda p: GM[p.year][p.q - 1] if p.actual else k.gross_margin[p.year - 2026], fmt=PCT, annual=None)
    rd = a.line("R&D", "% rev", lambda p: 0.14 if p.actual else 0.135, fmt=PCT, annual=None)
    sm = a.line("Sales & marketing", "% rev", lambda p: 0.11 if p.actual else 0.105, fmt=PCT, annual=None)
    ga = a.line("General & administrative", "% rev", lambda p: 0.06, fmt=PCT, annual=None)
    tax = a.line("Tax rate", "%", lambda p: 0.21, fmt=PCT, annual=None)
    a.section("Capital and working capital")
    capex = a.line("Capital expenditure", "% rev", lambda p: 0.07, fmt=PCT, annual=None)
    dep = a.line("Depreciation", "% PP&E/qtr", lambda p: 0.045, fmt=PCT, annual=None)
    dso = a.line("Receivable days", "days", lambda p: 52 if p.actual else 50, fmt="0", annual=None)
    dpo = a.line("Payable days", "days", lambda p: 41, fmt="0", annual=None)
    a.section("Valuation")
    a.ws.cell(a.row, 1, "Cost of capital (WACC)")
    a.ws.cell(a.row, 3, k.wacc).number_format = PCT
    a.ws.cell(a.row, 3).font = BLUE
    a.ws.cell(a.row, 3).fill = INPUT
    wacc_ref = f"Assumptions!$C${a.row}"
    a.row += 1
    a.ws.cell(a.row, 1, "Terminal growth")
    a.ws.cell(a.row, 3, k.terminal_growth).number_format = PCT
    a.ws.cell(a.row, 3).font = BLUE
    a.ws.cell(a.row, 3).fill = INPUT
    g_ref = f"Assumptions!$C${a.row}"
    a.row += 1
    a.ws.cell(a.row, 1, "Diluted shares")
    a.ws.cell(a.row, 2, "m").font = NOTE
    a.ws.cell(a.row, 3, 182.0).number_format = NUM
    a.ws.cell(a.row, 3).font = BLUE
    shares_ref = f"Assumptions!$C${a.row}"

    A = lambda p, r: ref("Assumptions", COL[(p.year, p.q)], r)  # noqa: E731

    # ------------------------------------------------------- income statement
    i = Sheet(wb, "Income Statement", "USD millions. Fiscal year ends 31 December.")
    i.section("Revenue build")
    rev_row = i.row
    i.row += 1  # revenue total goes here, written after its parts
    hw = i.line("Hardware revenue", "USDm", lambda p: f"={A(p, units)}*{A(p, asp)}", indent=1)
    i.line("Robots shipped", "k", lambda p: f"={A(p, units)}", fmt='#,##0.00', indent=2, memo=True)
    i.line("Average selling price", "USD k", lambda p: f"={A(p, asp)}", fmt='#,##0.0', indent=2, memo=True, annual=lambda p: f"={A(p, asp)}")
    svc = None
    if k.services:
        svc = i.line("Services revenue", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{hw}*{A(p, attach)}", indent=1)
    parts = [hw] + ([svc] if svc else [])
    i.row, after = rev_row, i.row
    i.line("Revenue", "USDm", lambda p: "=" + "+".join(f"{L(COL[(p.year, p.q)])}{r}" for r in parts), total=True)
    i.row = after
    i.line(
        "Revenue growth, year on year",
        "%",
        lambda p: None if p.year == 2024 else f"={L(COL[(p.year, p.q)])}{rev_row}/{L(COL[(p.year - 1, p.q)])}{rev_row}-1",
        fmt=PCT,
        memo=True,
        annual=lambda p: None if p.year == 2024 else f"={L(COL[(p.year, 0)])}{rev_row}/{L(COL[(p.year - 1, 0)])}{rev_row}-1",
    )
    i.section("Profit and loss")
    cogs = i.line("Cost of revenue", "USDm", lambda p: f"=-{L(COL[(p.year, p.q)])}{rev_row}*(1-{A(p, gm)})")
    gp = i.line("Gross profit", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{rev_row}+{L(COL[(p.year, p.q)])}{cogs}", total=True)
    i.line("Gross margin", "%", lambda p: f"={L(COL[(p.year, p.q)])}{gp}/{L(COL[(p.year, p.q)])}{rev_row}", fmt=PCT, memo=True, annual=lambda p: f"={L(COL[(p.year, 0)])}{gp}/{L(COL[(p.year, 0)])}{rev_row}")
    r1 = i.line("Research & development", "USDm", lambda p: f"=-{L(COL[(p.year, p.q)])}{rev_row}*{A(p, rd)}", indent=1)
    r2 = i.line("Sales & marketing", "USDm", lambda p: f"=-{L(COL[(p.year, p.q)])}{rev_row}*{A(p, sm)}", indent=1)
    r3 = i.line("General & administrative", "USDm", lambda p: f"=-{L(COL[(p.year, p.q)])}{rev_row}*{A(p, ga)}", indent=1)
    ebit = i.line("Operating income", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{gp}+SUM({L(COL[(p.year, p.q)])}{r1}:{L(COL[(p.year, p.q)])}{r3})", total=True)
    i.line("Operating margin", "%", lambda p: f"={L(COL[(p.year, p.q)])}{ebit}/{L(COL[(p.year, p.q)])}{rev_row}", fmt=PCT, memo=True, annual=lambda p: f"={L(COL[(p.year, 0)])}{ebit}/{L(COL[(p.year, 0)])}{rev_row}")
    taxr = i.line("Tax", "USDm", lambda p: f"=-MAX(0,{L(COL[(p.year, p.q)])}{ebit})*{A(p, tax)}")
    ni = i.line("Net income", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{ebit}+{L(COL[(p.year, p.q)])}{taxr}", total=True)
    i.line("Earnings per share", "USD", lambda p: f"={L(COL[(p.year, p.q)])}{ni}/{shares_ref}", fmt='0.00;(0.00)', annual="sum")

    IS = lambda p, r: ref("Income Statement", COL[(p.year, p.q)], r)  # noqa: E731

    # ------------------------------------------------------------- cash flow
    # Balance sheet rows are referenced before they exist; fix their numbers first.
    BS_CASH, BS_AR, BS_PPE, BS_AP = 7, 8, 9, 13

    cf = Sheet(wb, "Cash Flow", "USD millions.")
    cf.section("Operating")
    cf_ni = cf.line("Net income", "USDm", lambda p: f"={IS(p, ni)}")
    cf_dep = cf.line(
        "Depreciation",
        "USDm",
        lambda p: f"={ref('Balance Sheet', COL[(prev_quarter(p).year, prev_quarter(p).q)], BS_PPE)}*{A(p, dep)}" if prev_quarter(p) else 9.5,
    )
    cf_wc = cf.line(
        "Change in working capital",
        "USDm",
        lambda p: (
            f"=-({ref('Balance Sheet', COL[(p.year, p.q)], BS_AR)}-{ref('Balance Sheet', COL[(prev_quarter(p).year, prev_quarter(p).q)], BS_AR)})"
            f"+({ref('Balance Sheet', COL[(p.year, p.q)], BS_AP)}-{ref('Balance Sheet', COL[(prev_quarter(p).year, prev_quarter(p).q)], BS_AP)})"
        )
        if prev_quarter(p)
        else -4.0,
    )
    cfo = cf.line("Cash from operations", "USDm", lambda p: f"=SUM({L(COL[(p.year, p.q)])}{cf_ni}:{L(COL[(p.year, p.q)])}{cf_wc})", total=True)
    cf.section("Investing")
    cf_capex = cf.line("Capital expenditure", "USDm", lambda p: f"=-{IS(p, rev_row)}*{A(p, capex)}")
    fcf = cf.line("Free cash flow", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{cfo}+{L(COL[(p.year, p.q)])}{cf_capex}", total=True)
    cf.line("Free cash flow margin", "%", lambda p: f"={L(COL[(p.year, p.q)])}{fcf}/{IS(p, rev_row)}", fmt=PCT, memo=True, annual=lambda p: f"={L(COL[(p.year, 0)])}{fcf}/{IS(p, rev_row)}")

    # ---------------------------------------------------------- balance sheet
    b = Sheet(wb, "Balance Sheet", "USD millions, at period end. Annual columns show the year end.")
    b.section("Assets")
    assert b.row == BS_CASH
    cash = b.line(
        "Cash",
        "USDm",
        lambda p: f"={L(COL[(prev_quarter(p).year, prev_quarter(p).q)])}{BS_CASH}+{ref('Cash Flow', COL[(p.year, p.q)], fcf)}" if prev_quarter(p) else 212.0,
        annual="last",
    )
    ar = b.line("Receivables", "USDm", lambda p: f"={IS(p, rev_row)}*{A(p, dso)}/91", annual="last")
    ppe = b.line(
        "Property, plant & equipment",
        "USDm",
        lambda p: f"={L(COL[(prev_quarter(p).year, prev_quarter(p).q)])}{BS_PPE}-{ref('Cash Flow', COL[(p.year, p.q)], cf_capex)}-{ref('Cash Flow', COL[(p.year, p.q)], cf_dep)}" if prev_quarter(p) else 211.0,
        annual="last",
    )
    ta = b.line("Total assets", "USDm", lambda p: f"=SUM({L(COL[(p.year, p.q)])}{cash}:{L(COL[(p.year, p.q)])}{ppe})", total=True, annual="last")
    b.section("Liabilities and equity")
    assert b.row == BS_AP
    ap = b.line("Payables", "USDm", lambda p: f"=-{IS(p, cogs)}*{A(p, dpo)}/91", annual="last")
    eq = b.line("Equity", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{ta}-{L(COL[(p.year, p.q)])}{ap}", annual="last")
    tle = b.line("Total liabilities and equity", "USDm", lambda p: f"={L(COL[(p.year, p.q)])}{ap}+{L(COL[(p.year, p.q)])}{eq}", total=True, annual="last")
    assert (cash, ar, ppe) == (BS_CASH, BS_AR, BS_PPE)

    # -------------------------------------------------------------- valuation
    v = wb.create_sheet("Valuation")
    v["A1"] = "Northwind Robotics — Valuation"
    v["A1"].font = TITLE
    v["A2"] = "Discounted cash flow on the forecast years, terminal value on FY27."
    v["A2"].font = NOTE
    v.column_dimensions["A"].width = 40
    v.column_dimensions["B"].width = 14
    fy26, fy27 = COL[(2026, 0)], COL[(2027, 0)]
    rows = [
        ("Free cash flow FY26E", f"={ref('Cash Flow', fy26, fcf)}", NUM, GREEN),
        ("Free cash flow FY27E", f"={ref('Cash Flow', fy27, fcf)}", NUM, GREEN),
        ("Cost of capital", f"={wacc_ref}", PCT, GREEN),
        ("Terminal growth", f"={g_ref}", PCT, GREEN),
        ("Terminal value", "=B4*(1+B6)/(B5-B6)", NUM, BLACK),
        ("Enterprise value", "=B3/(1+B5)+(B4+B7)/(1+B5)^2", NUM, BLACK),
        ("Net cash", f"={ref('Balance Sheet', COL[(2025, 0)], BS_CASH)}", NUM, GREEN),
        ("Equity value", "=B8+B9", NUM, BOLD),
        ("Diluted shares", f"={shares_ref}", NUM, GREEN),
        ("Value per share (USD)", "=B10/B11", '0.00', BOLD),
        ("EV / FY26E operating income", f"=B8/{ref('Income Statement', fy26, ebit)}", MULT, BLACK),
    ]
    for n, (label, formula, fmt, font) in enumerate(rows, start=3):
        v.cell(n, 1, label)
        c = v.cell(n, 2, formula)
        c.number_format = fmt
        c.font = font
    for r in (10, 12):
        v.cell(r, 1).font = BOLD
        v.cell(r, 2).border = TOP

    # ----------------------------------------------------------------- checks
    ch = Sheet(wb, "Checks", "Every check should read zero.")
    ch.section("Balance sheet")
    ch.line("Assets less liabilities and equity", "USDm", lambda p: f"={ref('Balance Sheet', COL[(p.year, p.q)], ta)}-{ref('Balance Sheet', COL[(p.year, p.q)], tle)}", annual="last")

    wb.move_sheet("Valuation", offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
    return wb


def save_atomically(wb: Workbook, path: str) -> None:
    """Write to a temporary file and rename it over the target, as editors do."""
    folder = os.path.dirname(os.path.abspath(path))
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=folder)
    os.close(fd)
    wb.save(tmp)
    os.replace(tmp, path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "demo/northwind.xlsx"
    save_atomically(build(Knobs()), out)
    print(f"wrote {out}")
